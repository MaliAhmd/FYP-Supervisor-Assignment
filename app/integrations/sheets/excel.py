"""Standalone database snapshot export; never edits an existing workbook."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from app.integrations.sheets.register import HEADERS, register_rows


def export_excel(db, path, *, today=None):
    path = Path(path)
    if path.suffix.lower() != '.xlsx':
        raise ValueError('Expected an .xlsx path')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Progress'
    values = [HEADERS + ['Supervisor notes', 'Supervisor decision']]
    values.extend(row + ['', ''] for row in register_rows(db, today=today).values())
    for row_number, row in enumerate(values, 1):
        for col, value in enumerate(row, 1):
            if len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value):
                raise ValueError('Cell text is not representable in Excel')
            cell = sheet.cell(row_number, col, value)
            # Untrusted student names/criteria must remain literal, never formulas.
            cell.data_type = 's'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='24586B')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for column in 'ABCDEFGHIJKLMN':
        sheet.column_dimensions[column].width = 24
    # Exclusive creation avoids destroying supervisor edits in an old export.
    with path.open('xb') as stream:
        workbook.save(stream)
    workbook.close()
    return path
